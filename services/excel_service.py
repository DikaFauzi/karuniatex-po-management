from io import BytesIO
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="245E86")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="D9E7EF")

def _style_header(ws):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=THIN)

def _autowidth(ws):
    for col in ws.columns:
        max_len = 0
        for c in col:
            max_len = max(max_len, len(str(c.value or "")))
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 2, 10), 35)

def build_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "ORDERS"
    ws.append(["NO_PO","TANGGAL","CUSTOMER","ESTIMASI_HARI","STATUS","CATATAN"])
    ws.append(["PO-260925-001","25/09/2026","CONTOH CUSTOMER",21,"PROSES","Contoh catatan"])

    wi = wb.create_sheet("ITEMS")
    wi.append(["NO_PO","JENIS_KAIN","WARNA","QTY_KG","QTY_ROLL","HARGA_PER_KG"])
    wi.append(["PO-260925-001","COMBED 24S","HITAM",50,2,65000])

    wp = wb.create_sheet("PAYMENTS")
    wp.append(["NO_PO","NOMINAL_DP","TANGGAL_JAM_DP"])
    wp.append(["PO-260925-001",500000,"25/09/2026 10:30:00"])

    for s in [ws,wi,wp]:
        _style_header(s); _autowidth(s); s.freeze_panes="A2"

    out=BytesIO(); wb.save(out); out.seek(0)
    return out

def build_report(orders):
    wb=Workbook()

    ws=wb.active
    ws.title="RINGKASAN"
    total_value=sum(o["total"] for o in orders)
    total_dp=sum(o["total_dp"] for o in orders)
    total_sisa=sum(o["sisa"] for o in orders)
    total_kg=sum(o["total_qty_kg"] for o in orders)
    total_roll=sum(o["total_roll"] for o in orders)

    ws.append(["RINGKASAN LAPORAN PO","NILAI"])
    ws.append(["Total PO",len(orders)])
    ws.append(["Total Qty (KG)",total_kg])
    ws.append(["Total Estimasi Roll",total_roll])
    ws.append(["Total Nilai Pesanan",total_value])
    ws.append(["Total DP",total_dp])
    ws.append(["Total Sisa Tagihan",total_sisa])

    wr=wb.create_sheet("LAPORAN PO")
    wr.append(["NO PO","DIBUAT","TANGGAL PO","CUSTOMER","STATUS","TARGET","TOTAL QTY KG","ESTIMASI ROLL","TOTAL PESANAN","TOTAL DP","DP TERAKHIR","SISA","JUMLAH ITEM","JUMLAH FOTO","CATATAN"])
    for o in orders:
        wr.append([o["no_po"],o["created_at"],o["tanggal"],o["customer"],o["status"],o["target"],o["total_qty_kg"],o["total_roll"],o["total"],o["total_dp"],o["dp_terakhir"],o["sisa"],len(o["items"]),len(o["images"]),o["catatan"]])

    wi=wb.create_sheet("DETAIL ITEM")
    wi.append(["NO PO","CUSTOMER","JENIS KAIN","WARNA","QTY KG","QTY ROLL","HARGA/KG","TOTAL ITEM"])
    for o in orders:
        for it in o["items"]:
            wi.append([o["no_po"],o["customer"],it["jenis_kain"],it["warna"],it["qty_kg"],it["qty_roll"],it["harga_per_kg"],it["qty_kg"]*it["harga_per_kg"]])

    wp=wb.create_sheet("RIWAYAT DP")
    wp.append(["NO PO","CUSTOMER","NOMINAL DP","TANGGAL / JAM DP"])
    for o in orders:
        for p in o["payments"]:
            wp.append([o["no_po"],o["customer"],p["nominal"],p["created_at"]])

    currency_fmt='"Rp" #,##0'
    datetime_fmt='dd/mm/yyyy hh:mm:ss'
    date_fmt='dd/mm/yyyy'

    for s in wb.worksheets:
        _style_header(s)
        s.freeze_panes="A2"
        s.auto_filter.ref=s.dimensions
        s.sheet_view.showGridLines=False
        for row in s.iter_rows(min_row=2):
            for cell in row:
                cell.border=Border(bottom=THIN)
                cell.alignment=Alignment(vertical="top")
        _autowidth(s)

    for cell in ("B5","B6","B7"):
        ws[cell].number_format=currency_fmt
    ws["B3"].number_format='#,##0.00'
    ws["B4"].number_format='#,##0.00'

    for row in range(2,wr.max_row+1):
        wr.cell(row,2).number_format=datetime_fmt
        wr.cell(row,3).number_format=date_fmt
        wr.cell(row,6).number_format=date_fmt
        wr.cell(row,7).number_format='#,##0.00'
        wr.cell(row,8).number_format='#,##0.00'
        for col in (9,10,11,12):
            wr.cell(row,col).number_format=currency_fmt

    for row in range(2,wi.max_row+1):
        wi.cell(row,5).number_format='#,##0.00'
        wi.cell(row,6).number_format='#,##0.00'
        wi.cell(row,7).number_format=currency_fmt
        wi.cell(row,8).number_format=currency_fmt

    for row in range(2,wp.max_row+1):
        wp.cell(row,3).number_format=currency_fmt
        wp.cell(row,4).number_format=datetime_fmt

    out=BytesIO()
    wb.save(out)
    out.seek(0)
    return out

def parse_import(fileobj):
    wb=load_workbook(fileobj, data_only=True)
    required={"ORDERS","ITEMS","PAYMENTS"}
    if not required.issubset(set(wb.sheetnames)):
        raise ValueError("Sheet wajib: ORDERS, ITEMS, PAYMENTS")

    def rows(ws):
        vals=list(ws.iter_rows(values_only=True))
        headers=[str(x or "").strip() for x in vals[0]]
        return [dict(zip(headers,r)) for r in vals[1:] if any(v not in (None,"") for v in r)]

    return {
      "orders": rows(wb["ORDERS"]),
      "items": rows(wb["ITEMS"]),
      "payments": rows(wb["PAYMENTS"])
    }
